#!/usr/bin/env python3
"""
scripts/extraer_resumenes.py
T3 - Extracción de resumen descriptivo por instrumento desde Ficha_Tecnica.docx.
Genera data/resumenes.csv con: family_id, resumen, piezas_usadas, origen, revisar.
"""

import os
import csv
import docx
import re

REPO_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BASE_DIR = "/Users/emorosini/Downloads/Instrumentos/BASE_INSTRUMENTAL_FINAL"
CATALOGO_CSV = os.path.join(BASE_DIR, "CATALOGO_MAESTRO.csv")
OUTPUT_CSV = os.path.join(REPO_DIR, "data", "resumenes.csv")

SYNONYMS = {
    "tipo": ["tipo de instrumento", "formato", "informante"],
    "items": ["número de ítems", "numero de items", "ítems", "items", "n.º de ítems", "n.o de ítems"],
    "respuesta": ["formato de respuesta", "respuesta", "escala/formato de respuesta"],
    "dimensiones": ["dimensiones/subescalas", "dimensiones declaradas", "dimensiones", "estructura / subescalas"],
    "tiempo": ["tiempo estimado", "tiempo estimado de aplicación", "tiempo de aplicación"],
    "administracion": ["administración", "administracion", "modalidad de administración", "modalidad de administracion", "tipo de administración"]
}

def is_valid(val):
    if not val:
        return False
    v = val.strip().lower()
    if v.startswith("no verificado") or v.startswith("no fijado") or v.startswith("no especificado") or v.startswith("no informado"):
        return False
    if "no se encontró" in v or "no se encontro" in v or "no se localizó" in v or "no se localizo" in v:
        return False
    if "consultar manual" in v or "no verificada" in v or "no evaluada" in v:
        return False
    return True

def clean_items(text):
    t = text.strip()
    m_tot = re.search(r"total[:\s]+(\d+)\s*ítems", t, re.I)
    if m_tot:
        return f"{m_tot.group(1)} ítems"
    m_num = re.search(r"^(\d+)\s*(?:a\s*\d+\s*)?ítems", t, re.I)
    if m_num:
        return f"{m_num.group(0).strip()}".rstrip(".")
    clean = re.sub(r"\s*\(.*?\)", "", t).strip().rstrip(".")
    if len(clean) > 30:
        m = re.search(r"(\d+)\s*ítems", t, re.I)
        if m:
            return m.group(0)
    return clean

def clean_respuesta(text):
    t = text.strip()
    m_lik = re.search(r"(?:escala\s+)?likert\s+de\s+(\d+)\s+puntos", t, re.I)
    if m_lik:
        return f"Likert de {m_lik.group(1)} puntos"
    m_esc = re.search(r"escala\s+de\s+(\d+)\s+puntos", t, re.I)
    if m_esc:
        return f"Escala de {m_esc.group(1)} puntos"
    if re.search(r"dicot[oó]mic|verdadero/falso|s[ií]/no", t, re.I):
        return "Dicotómica (Sí/No)"
    clean = re.sub(r"\s*\(.*?\)", "", t).strip().rstrip(".")
    if len(clean) > 35:
        clean = clean[:35].rsplit(" ", 1)[0]
    return clean

def clean_dimensiones(text):
    t = text.strip()
    if re.search(r"^unidimensional", t, re.I):
        return "Unidimensional"
    m = re.search(r"(\d+)\s*(dimensiones|factores|escalas|subescalas|dominios|clústeres|componentes|facetas)", t, re.I)
    if m:
        word = m.group(2).lower()
        if word in ["factores", "escalas", "subescalas", "dominios", "clústeres", "componentes", "facetas"]:
            return f"{m.group(1)} dimensiones"
        return f"{m.group(1)} {word}"
    clean = re.sub(r"\s*\(.*?\)", "", t).strip().rstrip(".")
    if len(clean) > 30:
        clean = clean[:30].rsplit(" ", 1)[0]
    return clean

def clean_tiempo(text):
    t = text.strip()
    m = re.search(r"(\d+(?:[-–]\d+)?)\s*(?:minutos|min)", t, re.I)
    if m:
        return f"{m.group(1)} min"
    clean = re.sub(r"\s*\(.*?\)", "", t).strip().rstrip(".")
    return clean[:25]

def clean_tipo(text):
    t = text.strip()
    clean = re.sub(r"\s*\(.*?\)", "", t).strip().rstrip(".")
    if len(clean) > 35:
        clean = clean[:35].rsplit(" ", 1)[0]
    return clean

def clean_admin(text):
    t = text.strip()
    clean = re.sub(r"\s*\(.*?\)", "", t).strip().rstrip(".")
    if ";" in clean:
        clean = clean.split(";")[0].strip()
    return clean[:30]

def main():
    print(f"Leyendo catálogo maestro desde: {CATALOGO_CSV}")
    with open(CATALOGO_CSV, "r", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    print(f"Total familias: {len(rows)}")

    output_rows = []
    con_resumen = 0
    mismatches = 0

    for r in rows:
        fid = r["family_id"].strip()
        sigla = r["sigla"].strip()
        carpeta = r["carpeta"].strip()
        p = os.path.join(BASE_DIR, carpeta, "Ficha_Tecnica.docx")
        if not os.path.exists(p):
            raise FileNotFoundError(f"Ficha técnica docx no encontrada: {p}")

        doc = docx.Document(p)
        raw_text = ""
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    raw_text += " " + cell.text
        for p_elem in doc.paragraphs:
            raw_text += " " + p_elem.text

        raw_fields = {}
        if len(doc.tables) > 0:
            for t in doc.tables:
                for row in t.rows:
                    if len(row.cells) >= 2:
                        raw_k = row.cells[0].text.strip().lower()
                        k = raw_k.split(".", 1)[-1].strip()
                        v = row.cells[1].text.strip().replace("\n", " ")
                        for cat, syns in SYNONYMS.items():
                            if cat not in raw_fields and (k in syns or raw_k in syns):
                                if is_valid(v):
                                    raw_fields[cat] = v
        else:
            for p_elem in doc.paragraphs:
                text = p_elem.text.strip()
                if ":" in text:
                    raw_k, v = text.split(":", 1)
                    raw_k = raw_k.strip().lower()
                    v = v.strip().replace("\n", " ")
                    for cat, syns in SYNONYMS.items():
                        if cat not in raw_fields and raw_k in syns:
                            if is_valid(v):
                                raw_fields[cat] = v

        pieces = []
        # Orden estricto: tipo, items, respuesta, dimensiones, tiempo, administracion
        if "tipo" in raw_fields:
            c_t = clean_tipo(raw_fields["tipo"])
            if c_t: pieces.append(c_t)
        if "items" in raw_fields:
            c_i = clean_items(raw_fields["items"])
            if c_i: pieces.append(c_i)
        if "respuesta" in raw_fields:
            c_r = clean_respuesta(raw_fields["respuesta"])
            if c_r: pieces.append(c_r)
        if "dimensiones" in raw_fields:
            c_d = clean_dimensiones(raw_fields["dimensiones"])
            if c_d: pieces.append(c_d)
        if "tiempo" in raw_fields:
            c_tm = clean_tiempo(raw_fields["tiempo"])
            if c_tm: pieces.append(c_tm)
        if "administracion" in raw_fields:
            c_a = clean_admin(raw_fields["administracion"])
            if c_a and not any(c_a.lower() in p_piece.lower() or p_piece.lower() in c_a.lower() for p_piece in pieces):
                pieces.append(c_a)

        if len(pieces) >= 2:
            resumen = " · ".join(pieces)
            con_resumen += 1
            # Validación de seguridad numérica
            nums_in_res = re.findall(r"\b\d+\b", resumen)
            for num in nums_in_res:
                if not re.search(r"\b" + num + r"\b", raw_text):
                    print(f"  [ALERTA NÚMERO] {sigla}: '{num}' no aparece en el texto original")
                    mismatches += 1
        else:
            resumen = ""

        output_rows.append({
            "family_id": fid,
            "resumen": resumen,
            "piezas_usadas": len(pieces),
            "origen": "regla",
            "revisar": 0
        })

    os.makedirs(os.path.dirname(OUTPUT_CSV), exist_ok=True)
    with open(OUTPUT_CSV, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=["family_id", "resumen", "piezas_usadas", "origen", "revisar"],
            lineterminator="\n"
        )
        writer.writeheader()
        writer.writerows(output_rows)

    print(f"Guardado: {OUTPUT_CSV} ({len(output_rows)} filas)")
    print(f"Familias con resumen (>= 2 piezas): {con_resumen} / {len(rows)} ({con_resumen/len(rows)*100:.1f}%)")
    print(f"Inconsistencias numéricas encontradas: {mismatches}")
    if mismatches == 0:
        print("✓ VALIDACIÓN EXITOSA: 0 números inventados.")

if __name__ == "__main__":
    main()
